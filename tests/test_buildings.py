import io
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from schedule import (
    DEFAULT_BUILDINGS,
    expand_plan_buildings,
    merge_buildings,
    parse_book,
    plan,
    read_buildings,
    resolve_building,
)

FIX = Path(__file__).parent / "fixtures"


class BuildingTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.main = parse_book((FIX / "main.xlsx").read_bytes())
        cls.en1 = parse_book((FIX / "en1.xlsx").read_bytes(), True)
        cls.en2 = parse_book((FIX / "en2.xlsx").read_bytes(), True)

    def test_real_main_legend_has_house_numbers(self):
        self.assertEqual(self.main["buildings"], DEFAULT_BUILDINGS)
        self.assertEqual(self.main["buildings"]["Л"], "Львовская, 1В")

    def test_both_english_legend_layouts_are_read(self):
        for book in [self.en1, self.en2]:
            self.assertEqual(book["buildings"]["Л"], "Львовская")
            self.assertEqual(book["buildings"]["БП"], "Большая Печёрская")
            self.assertEqual(len(book["buildings"]), 6)

    def test_main_and_english_use_full_address(self):
        group = next(g for g in self.main["groups"].values() if g["name"] == "26КНТ-6")
        items = [item for item in plan(self.main, group) if item["date"] == "2026-10-21"]
        self.assertTrue(items)
        self.assertTrue(all(item["building"] == "Львовская, 1В" for item in items))
        self.assertTrue(all("Корпус: Львовская, 1В" in item["text"] for item in items))
        group = next(g for g in self.en1["en"].values() if g["num"] == 6 and g["level"] == "ОСНОВНОЙ" and g["code"] == "КНТ")
        items = plan(self.main, {"rows": []}, self.en1, [group["id"]])
        self.assertTrue(any(item.get("building") == "Львовская, 1В" for item in items))
        self.assertFalse(any(item.get("building") == "Л" for item in items))

    def test_source_address_overrides_default_but_short_english_does_not(self):
        buildings = merge_buildings(DEFAULT_BUILDINGS, {"Л": "Львовская, 99"}, self.en1["buildings"])
        self.assertEqual(buildings["Л"], "Львовская, 99")

    def test_abbreviation_variants_and_changed_room_place_are_resolved(self):
        for code in ["Б.Покр", "БПокр", "БПок"]:
            self.assertEqual(resolve_building(code, DEFAULT_BUILDINGS), "Большая Покровская, 22")
        for code in ["Б.Печ", "БПеч"]:
            self.assertEqual(resolve_building(code, DEFAULT_BUILDINGS), "Большая Печёрская, 25/12")
        self.assertEqual(resolve_building("Львовская", DEFAULT_BUILDINGS), "Львовская, 1В")
        self.assertEqual(resolve_building("P", DEFAULT_BUILDINGS), "Родионова, 136")
        self.assertEqual(resolve_building("online", DEFAULT_BUILDINGS), "online")
        self.assertEqual(resolve_building("Алексеевская, 6/16", DEFAULT_BUILDINGS), "Алексеевская, 6/16")

    def test_legacy_snapshot_is_copied_and_expanded(self):
        old = [{"date": "2026-10-21", "time": "11:10-12:30", "text": "Алгебра\nАуд.: 323\nКорпус: Л"}]
        new = expand_plan_buildings(old)
        self.assertEqual(new[0]["building"], "Львовская, 1В")
        self.assertIn("Корпус: Львовская, 1В", new[0]["text"])
        self.assertTrue(old[0]["text"].endswith("Корпус: Л"))
        self.assertNotIn("building", old[0])

    def test_old_cached_book_uses_legend_fallback(self):
        book = dict(self.main)
        book.pop("buildings")
        group = next(g for g in book["groups"].values() if g["name"] == "26КНТ-6")
        items = [item for item in plan(book, group) if item["date"] == "2026-10-21"]
        self.assertTrue(all(item["building"] == "Львовская, 1В" for item in items))

    def test_custom_source_legend_overrides_fallback(self):
        workbook = load_workbook(FIX / "main.xlsx", data_only=True)
        for sheet in workbook:
            for row in sheet:
                for cell in row:
                    if cell.value == "корпус на ул. Львовской, 1В":
                        cell.value = "корпус на ул. Львовской, 99"
        output = io.BytesIO()
        workbook.save(output)
        workbook.close()
        book = parse_book(output.getvalue())
        group = next(g for g in book["groups"].values() if g["name"] == "26КНТ-6")
        items = [item for item in plan(book, group) if item["date"] == "2026-10-21"]
        self.assertTrue(all(item["building"] == "Львовская, 99" for item in items))

    def test_unknown_legend_is_preserved_without_guessing(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet["A1"] = "ОБОЗНАЧЕНИЯ КОРПУСОВ:\nНовая - Н / Другой корпус - Д"
        sheet["A2"] = "корпус на ул. Тестовой, 7"
        sheet["B2"] = "Т"
        buildings = read_buildings(workbook)
        workbook.close()
        self.assertEqual(buildings, {"Н": "Новая", "Д": "Другой корпус", "Т": "Тестовой, 7"})
        self.assertEqual(resolve_building("Н", buildings), "Новая")
        self.assertEqual(resolve_building("нет в легенде", buildings), "нет в легенде")


if __name__ == "__main__":
    unittest.main()
