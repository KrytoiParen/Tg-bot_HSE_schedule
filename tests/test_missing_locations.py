import tempfile
import unittest
from datetime import date
from pathlib import Path
from unittest.mock import AsyncMock, patch

from openpyxl import Workbook

import database
import schedule
from messages import _lesson_cells, format_schedule

FIX = Path(__file__).parent / "fixtures"


class MissingLocationTests(unittest.TestCase):
    day = date(2026, 10, 21)

    def sheet(self, headers=(), values=()):
        book = Workbook()
        self.addCleanup(book.close)
        sheet = book.active
        sheet.title = "1 курс"
        sheet["E1"] = "МАГИСТРАТУРА - 1 курс, 1 модуль (01.09. - 25.10.)"
        sheet.merge_cells("E1:O1")
        for address, value in (
            ("B4", "День"), ("C4", "Время"), ("E4", "26 УПРБ-1"),
            ("B6", "Ср"), ("C6", "18:10-19:30"), ("E6", "Предмет\nПреподаватель"),
            *headers, *values,
        ):
            sheet[address] = value
        return sheet

    def group(self, sheet):
        return next(iter(schedule.read_groups(sheet, 1, 2026).values()))

    def html(self, group):
        items = schedule.plan(dict(year=2026, weeks={}), group)
        return format_schedule(items, self.day, group["name"])["html"]

    def test_real_law_first_course_keeps_lessons_with_explicit_missing_location(self):
        book = schedule.parse_book((FIX / "law.xlsx").read_bytes())
        groups = [g for g in book["groups"].values() if g["course"] == 1]
        self.assertEqual({g["name"] for g in groups}, {"26 УПРБ-1", "26 УПРБ-2"})
        for group in groups:
            with self.subTest(group=group["name"]):
                self.assertTrue(all(not r["room"] and not r["place"] for r in group["rows"]))
                items = schedule.plan(book, group)
                today = [v for v in items if v["date"] == self.day.isoformat()]
                self.assertEqual(len(today), 2)
                self.assertEqual({v["time"] for v in today}, {"18:10-19:30", "19:40-21:00"})
                html = format_schedule(items, self.day, group["name"])["html"]
                self.assertIn("<table", html)
                self.assertEqual(html.count("Аудитория не указана"), 2)
                self.assertEqual(html.count("Корпус не указан"), 2)
                self.assertNotIn("занятий на этот день нет", html)

    def test_absent_location_columns_do_not_pick_a_neighbouring_groups_subject(self):
        sheet = self.sheet(headers=(("G4", "26 УПРБ-2"),), values=(("G6", "Другой предмет"),))
        groups = schedule.read_groups(sheet, 1, 2026)
        self.assertEqual(len(groups), 2)
        for group in groups.values():
            self.assertEqual(group["rows"][0]["room"], "")
            self.assertEqual(group["rows"][0]["place"], "")
            html = self.html(group)
            self.assertIn("Аудитория не указана", html)
            self.assertIn("Корпус не указан", html)
        self.assertNotIn("Другой предмет", self.html(next(iter(groups.values()))))

    def test_only_one_location_column_preserves_the_available_field(self):
        for header, value, field, present, missing in (
            ("Аудитория", "402", "room", "<b>Ауд. 402</b>", "Корпус не указан"),
            ("Корпус", "К", "place", "Костина, 2Б", "Аудитория не указана"),
        ):
            with self.subTest(header=header):
                group = self.group(self.sheet(headers=(("F4", header),), values=(("F6", value),)))
                self.assertEqual(group["rows"][0][field], value)
                html = self.html(group)
                self.assertIn(present, html)
                self.assertIn(missing, html)

    def test_location_headers_can_be_swapped_and_separated_by_a_blank_column(self):
        sheet = self.sheet(
            headers=(("G4", " КОРПУС : "), ("H4", "Аудитория №")),
            values=(("G6", "К"), ("H6", "402")),
        )
        html = self.html(self.group(sheet))
        self.assertIn("<b>Ауд. 402</b><br>Костина, 2Б", html)
        self.assertNotIn("не указан", html)

    def test_location_headers_below_a_merged_group_header_are_read(self):
        sheet = self.sheet(headers=(("F5", "Ауд."), ("G5", "Корпус")),
                           values=(("F6", "402"), ("G6", "К")))
        sheet.merge_cells("E4:G4")
        html = self.html(self.group(sheet))
        self.assertIn("<b>Ауд. 402</b><br>Костина, 2Б", html)

    def test_shared_lecture_uses_the_last_groups_available_location_columns(self):
        for headers, values, room, place in (
            ((("F4", "Ауд."), ("J4", "Ауд.")), (("J6", "402"),), "402", ""),
            ((("F4", "Корпус"), ("J4", "Корпус")), (("J6", "К"),), "", "К"),
        ):
            with self.subTest(room=room, place=place):
                sheet = self.sheet(headers=(*headers, ("I4", "26 УПРБ-2")), values=values)
                sheet.merge_cells("E6:I6")
                for group in schedule.read_groups(sheet, 1, 2026).values():
                    row = group["rows"][0]
                    self.assertEqual((row["room"], row["place"]), (room, place))

    def test_merged_location_cell_is_not_duplicated_as_a_building(self):
        sheet = self.sheet(headers=(("F4", "Ауд."), ("G4", "Корпус")), values=(("F6", "402"),))
        sheet.merge_cells("F6:G6")
        row = self.group(sheet)["rows"][0]
        self.assertEqual((row["room"], row["place"]), ("402", ""))

    def test_merged_location_cell_preserves_an_address_or_online_as_a_place(self):
        for value in ("Алексеевская, 6/16,\n5 этаж\n(офис Е-Промо)", "online"):
            with self.subTest(value=value):
                sheet = self.sheet(headers=(("F4", "Ауд."), ("G4", "Корпус")), values=(("F6", value),))
                sheet.merge_cells("F6:G6")
                row = self.group(sheet)["rows"][0]
                self.assertEqual((row["room"], row["place"]), ("", value))

    def test_shared_lecture_does_not_treat_the_next_subject_as_a_location(self):
        sheet = self.sheet(
            headers=(("F4", "Ауд."), ("G4", "Корпус"), ("I4", "26 УПРБ-2"),
                     ("J4", "Ауд."), ("K4", "Корпус"), ("M4", "26 УПРБ-3")),
            values=(("L6", "Другой предмет"), ("M6", "Ещё один предмет")),
        )
        sheet.merge_cells("E6:K6")
        group = self.group(sheet)
        self.assertEqual((group["rows"][0]["room"], group["rows"][0]["place"]), ("", ""))
        sheet["L6"] = "999"
        group = self.group(sheet)
        self.assertEqual((group["rows"][0]["room"], group["rows"][0]["place"]), ("999", ""))

    def test_room_and_building_values_are_not_borrowed_from_the_next_group(self):
        sheet = self.sheet(
            headers=(("I4", "26 УПРБ-2"), ("J4", "Ауд."), ("K4", "Корпус")),
            values=(("I6", "Другой предмет"), ("J6", "999"), ("K6", "Р")),
        )
        group = self.group(sheet)
        self.assertEqual((group["rows"][0]["room"], group["rows"][0]["place"]), ("", ""))
        self.assertNotIn("999", self.html(group))

    def test_blank_and_dash_fields_are_reported_as_missing(self):
        for suffix in ("", "\nАуд.: \nКорпус: ", "\nАуд.: —\nКорпус: -"):
            with self.subTest(suffix=suffix):
                subject, location = _lesson_cells(dict(text="Предмет\nПреподаватель" + suffix))
                self.assertIn("Предмет", subject)
                self.assertIn("Преподаватель", subject)
                self.assertEqual(location, "Аудитория не указана<br>Корпус не указан")

    def test_online_location_does_not_request_a_physical_room_or_building(self):
        for suffix in (
            "Корпус: online\nасинхронно", "Корпус: Онлайн", "Ауд.: online", "Ауд.: дистанционно",
            "Корпус: Дистанционное обучение", "Ауд.: on-line", "Ауд.: online\nасинхронно",
            "Ауд.: online\nасинхронно\nКорпус: online\nасинхронно",
        ):
            with self.subTest(suffix=suffix):
                _, location = _lesson_cells(dict(text="Предмет\n" + suffix))
                self.assertNotIn("не указан", location)
                self.assertNotIn("Ауд.", location)
                self.assertTrue(location)
                self.assertLessEqual(location.count("online"), 1)
                self.assertLessEqual(location.count("асинхронно"), 1)

    def test_missing_fields_in_a_legacy_row_do_not_remove_the_lesson(self):
        group = self.group(self.sheet())
        for row in group["rows"]:
            row.pop("room")
            row.pop("place")
        html = self.html(group)
        self.assertIn("Предмет", html)
        self.assertIn("Аудитория не указана", html)
        self.assertIn("Корпус не указан", html)


class LocationCacheTests(unittest.IsolatedAsyncioTestCase):
    async def test_previous_parser_cache_is_rebuilt_without_changing_group_ids(self):
        data = (FIX / "law.xlsx").read_bytes()
        previous = schedule.parse_book(data)
        previous["parser_version"] = 2
        with tempfile.TemporaryDirectory() as temp, \
                patch.object(database, "DB_PATH", Path(temp) / "bot.db"), \
                patch.object(schedule, "download", new=AsyncMock(return_value=data)) as download:
            database.create_database()
            database.save_cache("book:law", previous)
            book = await schedule.get_book("law")
            download.assert_awaited_once()
            self.assertEqual(book["parser_version"], schedule.PARSER_VERSION)
            self.assertEqual(set(book["groups"]), set(previous["groups"]))
            self.assertEqual(database.get_cache("book:law"), book)


if __name__ == "__main__":
    unittest.main()
