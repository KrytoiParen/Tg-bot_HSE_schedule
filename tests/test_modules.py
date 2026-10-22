import asyncio
import copy
import io
import re
import tempfile
import unittest
from datetime import date, datetime
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import AsyncMock, patch

from openpyxl import load_workbook

import database
import schedule
import utils
from config import TZ
from handlers import user as dialogs
from messages import format_schedule
from sources import ParseError, SourceUnavailable, discover_modules

FIX = Path(__file__).parent / "fixtures"
PAGE_URL = "https://nnov.hse.ru/uch/schedule"


def module_page(second=True):
    return """<h2>I модуль</h2><h4>Расписание учебных занятий:</h4>
        <a href="/first.pdf"><strong>очный</strong> формат обучения</a>
        <a href="/part-time.pdf">очно-заочный формат обучения</a>
        <h4>Расписание сессии:</h4><a href="/exam.pdf">очный формат обучения</a>
        <h2>II модуль</h2><h4>Расписание учебных занятий:</h4>
        <a href="SECOND">очный формат обучения</a>""".replace("SECOND", "/second.pdf" if second else "/not-published")


def repeated_book(layout):
    book = load_workbook(FIX / "main.xlsx")
    for sheet in list(book):
        if sheet.title != "Календарь":
            book.remove(sheet)
    first = book.create_sheet("1 курс")

    def block(sheet, row, col, module, title):
        sheet.cell(row, col, f"БАКАЛАВРИАТ - 1 курс, {module} модуль " +
                   ("(01.09.-25.10.)" if module == 1 else "(01.11.-20.12.)"))
        for c, value in ((2, "День"), (3, "Время"), (col, "26МТ-1"), (col + 1, "Ауд."), (col + 2, "Корпус")):
            sheet.cell(row + 2, c, value)
        for c, value in ((2, "Пн"), (3, "08:00-09:20"), (col, title), (col + 1, "101"), (col + 2, "Л")):
            sheet.cell(row + 3, c, value)

    block(first, 1, 5, 1, "Алгебра первого модуля")
    if layout == "horizontal":
        block(first, 1, 9, 2, "Анализ второго модуля")
    elif layout == "vertical":
        block(first, 7, 5, 2, "Анализ второго модуля")
    else:
        second = book.create_sheet("1 курс II модуль")
        block(second, 1, 5, 2, "Анализ второго модуля")
    stream = io.BytesIO()
    book.save(stream)
    book.close()
    return stream.getvalue()


class ModuleParserTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.book = schedule.parse_book((FIX / "main.xlsx").read_bytes())

    def test_real_page_excludes_part_time_and_preserves_unpublished_modules(self):
        modules = discover_modules((FIX / "schedule.html").read_text(), PAGE_URL)
        self.assertEqual(set(modules), {"1", "2", "3", "4"})
        self.assertTrue(modules["1"]["url"].endswith("1200526994.pdf"))
        self.assertTrue(all(modules[n]["url"] is None for n in ("2", "3", "4")))

    def test_only_full_time_teaching_links_are_selected(self):
        self.assertEqual(discover_modules(module_page(), PAGE_URL), {
            "1": {"url": "https://nnov.hse.ru/first.pdf"},
            "2": {"url": "https://nnov.hse.ru/second.pdf"},
        })

    def test_missing_or_ambiguous_sources_fail_instead_of_guessing(self):
        for html in ("<html>Сервис недоступен</html>", module_page() +
                     '<a href="/another.pdf">очный формат обучения</a>'):
            with self.subTest(html=html), self.assertRaises(ParseError):
                discover_modules(html, PAGE_URL)

    def test_official_calendar_identifies_all_modules_and_sessions(self):
        days = self.book["calendar_days"]
        for stamp, module, kind in (
            ("2026-10-22", 1, "classes"), ("2026-10-26", 1, "session"),
            ("2026-11-02", 2, "classes"), ("2027-01-11", 3, "classes"),
            ("2027-04-05", 4, "classes"), ("2027-05-03", 4, "vacation"),
        ):
            with self.subTest(stamp=stamp):
                self.assertEqual(days[stamp], {"module": module, "kind": kind})

    def test_same_group_in_multiple_module_blocks_is_not_overwritten(self):
        for layout in ("horizontal", "vertical", "sheets"):
            with self.subTest(layout=layout):
                book = schedule.parse_book(repeated_book(layout))
                self.assertEqual(len(book["groups"]), 1)
                group = next(iter(book["groups"].values()))
                self.assertEqual(group["id"], schedule.key("1 курс", schedule.norm("26МТ-1")))
                items = schedule.plan(book, group)
                early = [v for v in items if v["date"] == "2026-09-07"]
                late = [v for v in items if v["date"] == "2026-11-02"]
                self.assertEqual(len(early), 1)
                self.assertEqual(len(late), 1)
                self.assertIn("Алгебра первого", early[0]["text"])
                self.assertIn("Анализ второго", late[0]["text"])

    def test_module_annotation_uses_header_number_instead_of_month(self):
        row = dict(day=3, time="08:00-09:20", text="Со 2 модуля Алгебра", room="101", place="Л",
                   start="2026-10-01", end="2026-11-30", mode="all", module=2)
        self.assertEqual(len(schedule.on_day(row, date(2026, 10, 22), 2026, {})), 1)
        self.assertEqual(schedule.on_day(dict(row, module=1), date(2026, 10, 22), 2026, {}), [])

    def test_overlapping_table_periods_do_not_mix_different_modules(self):
        book = schedule.parse_book(repeated_book("horizontal"))
        group = next(iter(book["groups"].values()))
        for row in group["rows"]:
            row.update(start="2026-09-01", end="2026-12-20")
        items = schedule.plan(book, group)
        self.assertTrue(all("Алгебра первого" not in v["text"] for v in items if v["date"] == "2026-11-02"))
        self.assertTrue(all("Анализ второго" not in v["text"] for v in items if v["date"] == "2026-10-19"))

    def test_bachelor_mathematics_first_course_has_both_groups_and_lessons(self):
        book = schedule.parse_book((FIX / "mathematics.xlsx").read_bytes())
        groups = [g for g in book["groups"].values() if g["code"] in schedule.CODES["Математика"] and g["course"] == 1]
        self.assertEqual({g["name"] for g in groups}, {"26 ФПМ-1", "26 ФПМ-2"})
        for group in groups:
            items = [v for v in schedule.plan(book, group) if v["date"] == "2026-10-19"]
            self.assertTrue(items)
            self.assertIn("<table", format_schedule(items, date(2026, 10, 19), group["name"])["html"])


class ModuleServiceTests(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.patches = [patch.object(database, "DB_PATH", Path(self.temp.name) / "bot.db"),
                        patch.object(schedule, "source_index", {}), patch.object(schedule, "refs", []),
                        patch.object(schedule, "lock", asyncio.Lock())]
        for item in self.patches:
            item.start()
        database.create_database()
        self.first = schedule.parse_book(repeated_book("sheets"))
        self.group = next(iter(self.first["groups"].values()))
        self.second = copy.deepcopy(self.first)
        self.ref = dict(id="stable-program", book="first-book", name="Математика", level="БАКАЛАВРИАТ",
                        faculty="Факультет математики", english=False)
        self.second_ref = dict(self.ref, book="second-book")
        self.user = dict(program=self.ref["id"], group_id=self.group["id"], en=[], snapshot=[])

    async def asyncTearDown(self):
        for item in reversed(self.patches):
            item.stop()
        self.temp.cleanup()

    def index(self, second=True):
        schedule.source_index = dict(year=2026, days=self.first["calendar_days"], modules={
            "1": dict(url="https://nnov.hse.ru/first.pdf", refs=[self.ref]),
            "2": dict(url="https://nnov.hse.ru/second.pdf" if second else None, refs=[self.second_ref] if second else []),
        })
        schedule.refs = [self.ref]

    async def test_index_discovers_pdfs_and_caches_calendar_without_part_time_requests(self):
        pdf = (FIX / "index.pdf").read_bytes()

        async def download(url, kind):
            return module_page().encode() if kind == "html" else pdf

        with patch.object(schedule, "download", new=AsyncMock(side_effect=download)) as fetch, \
                patch.object(schedule, "get_book", new=AsyncMock(return_value=self.first)), \
                patch.object(schedule, "datetime") as clock:
            clock.now.return_value = datetime(2026, 10, 22, tzinfo=TZ)
            await schedule.get_index()
        urls = [call.args[0] for call in fetch.await_args_list]
        self.assertEqual(urls, [PAGE_URL, "https://nnov.hse.ru/first.pdf", "https://nnov.hse.ru/second.pdf"])
        self.assertEqual(schedule.calendar_day(date(2027, 1, 11))["module"], 3)
        self.assertEqual(database.get_cache("source-index:v1"), schedule.source_index)

    async def test_same_subscription_selects_different_module_catalogs(self):
        self.index()
        database.save_cache("book:first-book", self.first)
        database.save_cache("book:second-book", self.second)
        before = copy.deepcopy(self.user)
        with patch.object(schedule, "download", new=AsyncMock()) as fetch:
            first, _, early_notice = await utils.daily_schedule(self.user, date(2026, 10, 19))
            second, _, late_notice = await utils.daily_schedule(self.user, date(2026, 11, 2))
        self.assertIsNone(early_notice)
        self.assertIsNone(late_notice)
        self.assertIn("Алгебра первого", next(v["text"] for v in first if v["date"] == "2026-10-19"))
        self.assertIn("Анализ второго", next(v["text"] for v in second if v["date"] == "2026-11-02"))
        self.assertEqual(schedule.get_program(self.user["program"], date(2026, 11, 2))["book"], "second-book")
        self.assertEqual(self.user, before)
        fetch.assert_not_awaited()

    async def test_unpublished_module_is_not_reported_as_no_lessons_or_old_module(self):
        self.index(second=False)
        items, name, notice = await utils.daily_schedule(self.user, date(2026, 11, 2))
        self.assertEqual(items, [])
        self.assertIn("модуля 2", notice)
        html = format_schedule(items, date(2026, 11, 2), name, notice=notice)["html"]
        self.assertIn("ещё не опубликовано", html)
        self.assertNotIn("занятий на этот день нет", html)

    async def test_next_academic_year_cannot_reuse_last_year_calendar(self):
        self.index()
        self.assertIn("нет актуального учебного календаря", schedule.source_notice(date(2027, 9, 1)))

    async def test_session_and_holidays_have_explicit_notices(self):
        self.index()
        self.assertIn("сессия", schedule.source_notice(date(2026, 10, 26)))
        self.assertIn("каникулы", schedule.source_notice(date(2027, 5, 3)))

    async def test_pdf_pointing_to_old_module_does_not_produce_a_false_empty_day(self):
        self.index()
        stale = copy.deepcopy(self.first)
        for group in stale["groups"].values():
            group["rows"] = [row for row in group["rows"] if row["module"] == 1]
            group["periods"] = [period for period in group["periods"] if period["module"] == 1]
        database.save_cache("book:second-book", stale)
        items, _, notice = await utils.daily_schedule(self.user, date(2026, 11, 2))
        self.assertEqual(items, [])
        self.assertIn("ещё не содержит расписание", notice)

    async def test_overlapping_dates_cannot_make_a_wrong_module_look_valid(self):
        self.index()
        stale = copy.deepcopy(self.first)
        for group in stale["groups"].values():
            group["periods"] = [dict(period, end="2026-12-20") for period in group["periods"] if period["module"] == 1]
        database.save_cache("book:second-book", stale)
        _, _, notice = await utils.daily_schedule(self.user, date(2026, 11, 2))
        self.assertIn("ещё не содержит расписание", notice)

    async def test_a_published_empty_group_is_distinguished_from_an_unpublished_module(self):
        self.index()
        empty = copy.deepcopy(self.second)
        empty["groups"][self.group["id"]]["rows"] = []
        database.save_cache("book:second-book", empty)
        items, _, notice = await utils.daily_schedule(self.user, date(2026, 11, 2))
        self.assertEqual(items, [])
        self.assertIsNone(notice)

    async def test_a_groups_shorter_teaching_period_does_not_mean_its_module_is_unpublished(self):
        self.index()
        book = copy.deepcopy(self.first)
        group = book["groups"][self.group["id"]]
        for period in group["periods"] + group["rows"]:
            if period["module"] == 1:
                period["end"] = "2026-10-18"
        database.save_cache("book:first-book", book)
        items, _, notice = await utils.daily_schedule(self.user, date(2026, 10, 19))
        self.assertIsNone(notice)
        self.assertFalse(any(item["date"] == "2026-10-19" for item in items))

    async def test_week_command_selects_each_days_module(self):
        self.index()
        for bid in ("first-book", "second-book"):
            database.save_cache("book:" + bid, self.first)
        database.save_user(9, self.user["program"], self.user["group_id"], [], "07:00", [])
        bot = SimpleNamespace(send_rich_message=AsyncMock(), send_message=AsyncMock())
        message = SimpleNamespace(chat=SimpleNamespace(id=9), bot=bot, text="/week")
        with patch.object(dialogs, "datetime") as clock:
            clock.now.return_value = datetime(2026, 10, 30, tzinfo=TZ)
            await dialogs.show(message)
        self.assertEqual(bot.send_rich_message.await_count, 7)
        html = "\n".join(call.kwargs["rich_message"].html for call in bot.send_rich_message.await_args_list)
        self.assertIn("сессия", html)
        self.assertIn("Анализ второго модуля", html)
        self.assertNotIn("Алгебра первого модуля", html)

    async def test_morning_queue_reports_unpublished_module_without_false_empty_schedule(self):
        self.index(second=False)
        database.save_user(9, self.user["program"], self.user["group_id"], [], "07:00", [])
        utils.morning_mail(datetime(2026, 11, 2, 7, 1, tzinfo=TZ))
        html = database.get_mail()[0]["text"]["html"]
        self.assertIn("ещё не опубликовано", html)
        self.assertNotIn("занятий на этот день нет", html)

    async def test_old_parsed_cache_is_upgraded_before_use(self):
        database.save_cache("book:old", dict(parser_version=1, groups={}))
        with patch.object(schedule, "download", new=AsyncMock(return_value=(FIX / "main.xlsx").read_bytes())) as fetch:
            book = await schedule.get_book("old")
        self.assertEqual(book["parser_version"], schedule.PARSER_VERSION)
        self.assertTrue(book["calendar_days"])
        fetch.assert_awaited_once()

    async def test_source_index_survives_restart(self):
        self.index()
        expected = copy.deepcopy(schedule.source_index)
        database.save_cache("source-index:v1", expected)
        database.save_cache("catalog", [self.ref])
        with patch.object(schedule, "http", None):
            try:
                await schedule.open_sources()
                self.assertEqual(schedule.source_index, expected)
                self.assertEqual(schedule.get_catalog(date(2026, 11, 2)), [self.second_ref])
            finally:
                await schedule.close_sources()

    async def test_legacy_catalog_is_not_used_before_module_calendar_is_verified(self):
        database.save_cache("catalog", [self.ref])
        with patch.object(schedule, "http", None):
            try:
                await schedule.open_sources()
                self.assertEqual(schedule.refs, [])
                self.assertIn("ещё не проверены", schedule.source_notice(date(2026, 11, 2)))
            finally:
                await schedule.close_sources()

    async def test_missing_english_cache_does_not_use_a_previous_module_snapshot(self):
        self.index()
        english_ref = dict(id="english", name="1 курс", book="english-two", english=True)
        schedule.source_index["modules"]["2"]["refs"].append(english_ref)
        database.save_cache("book:second-book", self.second)
        user = dict(self.user, en=["selected-english"], snapshot=[
            dict(date="2026-10-19", time="08:00-09:20", text="Старое расписание")])
        with self.assertRaises(SourceUnavailable):
            utils.cached_plan(user, date(2026, 11, 2))

    async def test_english_group_id_survives_another_module_sheet(self):
        self.index()
        workbook = load_workbook(FIX / "en1.xlsx")
        second = workbook.copy_worksheet(workbook["1 курс"])
        second.title = "1 курс II модуль"
        for row in second:
            for cell in row:
                if isinstance(cell.value, str) and "модул" in cell.value.lower():
                    cell.value = re.sub(r"1\s*модуль", "2 модуль", cell.value, flags=re.I)
                    cell.value = re.sub(r"\([^)]*\)", "(01.11.-20.12.)", cell.value)
        data = io.BytesIO()
        workbook.save(data)
        workbook.close()
        english = schedule.parse_book(data.getvalue(), True)
        original = schedule.parse_book((FIX / "en1.xlsx").read_bytes(), True)
        self.assertEqual(set(english["en"]), set(original["en"]))
        gid = next(iter(english["en"]))
        self.assertEqual({row["module"] for row in english["en"][gid]["rows"]}, {1, 2})
        schedule.validate_english(english, self.first, [gid], date(2026, 11, 2))

    async def test_pdf_failure_is_isolated_to_its_module(self):
        async def download(url, kind):
            if kind == "html":
                return module_page().encode()
            if url.endswith("second.pdf"):
                raise OSError("unavailable")
            return (FIX / "index.pdf").read_bytes()

        with patch.object(schedule, "download", new=AsyncMock(side_effect=download)), \
                patch.object(schedule, "get_book", new=AsyncMock(return_value=self.first)), \
                patch.object(schedule, "datetime") as clock:
            clock.now.return_value = datetime(2026, 10, 22, tzinfo=TZ)
            await schedule.get_index()
        self.assertTrue(schedule.get_catalog(date(2026, 10, 22)))
        with self.assertRaises(SourceUnavailable):
            schedule.get_catalog(date(2026, 11, 2))

    async def test_mathematics_first_course_is_available_through_real_dialog_handlers(self):
        book = schedule.parse_book((FIX / "mathematics.xlsx").read_bytes())
        schedule.refs = [self.ref]

        class State:
            def __init__(inner):
                inner.data = dict(step="program", nonce="initial", options=[dict(value=self.ref["id"])])

            async def get_data(inner):
                return inner.data.copy()

            async def update_data(inner, **kwargs):
                inner.data.update(kwargs)

            async def set_state(inner, state):
                inner.state = state

        state = State()
        message = SimpleNamespace(chat=SimpleNamespace(type="private"), answer=AsyncMock(), edit_reply_markup=AsyncMock())
        query = SimpleNamespace(message=message, data="pick:initial:0", answer=AsyncMock())
        with patch.object(schedule, "get_book", new=AsyncMock(return_value=book)):
            await dialogs.callback_query(query, state, None)
        self.assertIn(1, [option["value"] for option in state.data["options"]])
        first = next(i for i, option in enumerate(state.data["options"]) if option["value"] == 1)
        query.data = f"pick:{state.data['nonce']}:{first}"
        await dialogs.callback_query(query, state, None)
        text = message.answer.await_args.args[0]
        self.assertIn("26 ФПМ-1", text)
        self.assertIn("26 ФПМ-2", text)


if __name__ == "__main__":
    unittest.main()
