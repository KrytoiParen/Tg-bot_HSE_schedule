import asyncio
import copy
import io
import tempfile
import unittest
from datetime import date, datetime, timedelta
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import AsyncMock, patch

from openpyxl import load_workbook

import database
import schedule as source
import utils as bot_code
from handlers import user as user_code
from config import TZ
from utils import valid_time, changes, split
from schedule import catalog, parse_book, en_options, plan, on_day, date_rules, ParseError, current_english

FIX = Path(__file__).parent / "fixtures"


class ParserTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.refs = catalog((FIX / "index.pdf").read_bytes())
        cls.ref = next(r for r in cls.refs if r["name"] == "Компьютерные науки и технологии")
        cls.book = parse_book((FIX / "main.xlsx").read_bytes())
        cls.en1 = parse_book((FIX / "en1.xlsx").read_bytes(), True)
        cls.en2 = parse_book((FIX / "en2.xlsx").read_bytes(), True)
        cls.group = next(g for g in cls.book["groups"].values() if g["name"] == "26КНТ-1")

    def option(self, number, level, book=None, course=1, mode=None):
        group = dict(self.group, course=course)
        choices = en_options(book or self.en1, self.ref, group, mode)
        return next(e for e in choices if e["num"] == number and e["level"] == level and e["code"] == "КНТ")

    def test_catalog_is_links_with_faculties(self):
        self.assertEqual(len(self.refs), 27)
        self.assertEqual(len({r["book"] for r in self.refs}), 23)
        self.assertEqual(self.ref["faculty"], "Факультет информатики, математики и компьютерных наук")
        self.assertEqual(len([r for r in self.refs if r["english"]]), 2)

    def test_courses_and_groups(self):
        self.assertEqual(len(self.book["groups"]), 30)
        self.assertEqual(self.group["course"], 1)
        self.assertTrue(self.group["refs"])

    def test_english_merged_cell_uses_exact_line_and_level(self):
        a = self.option(8, "БАЗОВЫЙ")
        b = self.option(8, "ОСНОВНОЙ")
        d = date(2026, 10, 1)
        first = plan(self.book, {"rows": []}, self.en1, [a["id"]])
        second = plan(self.book, {"rows": []}, self.en1, [b["id"]])
        first = [v for v in first if v["date"] == d.isoformat()]
        second = [v for v in second if v["date"] == d.isoformat()]
        self.assertEqual(len(first), 2)
        self.assertTrue(all("Султанова А.В." in v["text"] and "Ауд.: 324" in v["text"] for v in first))
        self.assertTrue(all("Милащенко Н.Г." in v["text"] and "Ауд.: 309" in v["text"] for v in second))
        self.assertTrue(all("Галочкин" not in v["text"] and "Давыдова" not in v["text"] for v in first + second))

    def test_number_1_does_not_pick_11(self):
        e = self.option(1, "БАЗОВЫЙ")
        for row in e["rows"]:
            self.assertEqual(row["group"].split("(")[0].strip(), "гр 1")
        items = plan(self.book, {"rows": []}, self.en1, [e["id"]])
        self.assertTrue(any("Ватлецов" in v["text"] for v in items))
        self.assertFalse(any("Давыдова" in v["text"] for v in items))

    def test_faculty_and_program_filter(self):
        choices = en_options(self.en1, self.ref, self.group)
        self.assertFalse(any(e["code"] == "ТИДИ" for e in choices))
        self.assertFalse(any(e["faculty"] == "Факультет права" for e in choices))

    def test_shared_business_program_uses_shared_english_scope(self):
        ref = next(
            r
            for r in self.refs
            if r["name"] == "Международный бакалавриат по бизнесу и экономике" and r["faculty"] == "Факультет экономики"
        )
        choices = en_options(self.en1, ref, dict(self.group, code="МПБЭ"))
        self.assertTrue(any(e["code"] == "МПБЭ" for e in choices))
        self.assertTrue(any(not e["code"] and e["faculty"] == "Факультет менеджмента" for e in choices))

    def test_second_course_week_is_from_official_calendar(self):
        self.assertEqual(self.en2["year"], 2025)
        fixed = current_english(self.en2, self.book)
        self.assertEqual(fixed["year"], 2026)
        self.assertEqual(fixed["weeks"]["2026-09-14"], "upper")
        self.assertEqual(fixed["weeks"]["2026-09-21"], "lower")
        choices = en_options(self.en2, self.ref, dict(self.group, course=2))
        upper = next(e for e in choices if e["mode"] == "upper" and e["rows"])
        lower = next(e for e in choices if e["mode"] == "lower" and e["rows"])
        for e in [upper, lower]:
            items = plan(self.book, {"rows": []}, self.en2, [e["id"]])
            self.assertTrue(items)
            self.assertTrue(all(fixed["weeks"][v["date"]] == e["mode"] for v in items))

    def test_english_start_and_cancellation(self):
        e = self.option(6, "ОСНОВНОЙ")
        row = next(r for r in e["rows"] if r["day"] == 3 and r["time"] == "14:40-16:00")
        self.assertFalse(on_day(row, date(2026, 9, 17), 2026, self.en1["weeks"]))
        self.assertTrue(on_day(row, date(2026, 9, 24), 2026, self.en1["weeks"]))
        row = dict(
            day=0,
            time="09:30-10:50",
            text="с 14.09. Теория государства и права\nсеминар\n05.10. - отмена занятия",
            room="107",
            place="С",
            start="2026-09-01",
            end="2026-10-25",
            mode="all",
        )
        self.assertFalse(on_day(row, date(2026, 10, 5), 2026, {}))
        self.assertTrue(on_day(row, date(2026, 10, 12), 2026, {}))

    def test_date_ranges_and_lists(self):
        allowed, denied = date_rules("07.09.-19.10. Основы государственности\n12.10. отмена занятия", 2026)
        self.assertIn(date(2026, 10, 5), allowed)
        self.assertIn(date(2026, 10, 12), denied)
        allowed, _ = date_rules("07.09-19.10 Предмет", 2026)
        self.assertIn(date(2026, 10, 5), allowed)
        allowed, _ = date_rules("07.09., 21.09., 05.10., 19.10. Предмет", 2026)
        self.assertNotIn(date(2026, 10, 12), allowed)
        allowed, _ = date_rules("07.09, 14.09, с 05.10 по 19.10\nПредмет", 2026)
        self.assertNotIn(date(2026, 9, 28), allowed)
        self.assertIn(date(2026, 10, 12), allowed)

    def test_alternative_subjects_have_correct_rooms(self):
        row = dict(
            day=4,
            time="08:00-09:20",
            text="02.10., 16.10. Программирование\nЛупанова\n\n09.10., 23.10. Технологии\nУлитин",
            room="102\n\n\n203",
            place="Р",
            start="2026-09-01",
            end="2026-10-25",
            mode="all",
        )
        result = on_day(row, date(2026, 10, 9), 2026, {})
        self.assertEqual(len(result), 1)
        self.assertIn("Ауд.: 203", result[0]["text"])
        self.assertNotIn("Лупанова", result[0]["text"])

    def test_time_override_does_not_make_every_week(self):
        row = dict(
            day=1,
            time="11:10-12:30",
            text="20.10. в 11:10-11:50 Линейная алгебра - лекция - Хомицкий",
            room="302",
            place="БП",
            start="2026-09-01",
            end="2026-10-25",
            mode="all",
        )
        self.assertFalse(on_day(row, date(2026, 10, 13), 2026, {}))
        self.assertEqual(on_day(row, date(2026, 10, 20), 2026, {})[0]["time"], "11:10-11:50")

    def test_room_override(self):
        row = dict(
            day=3,
            time="08:00-09:20",
            text="Английский",
            teacher="Ватлецов С.Г. 24.09. в 324 (Львовская)",
            room="301",
            place="БП",
            start="2026-09-14",
            end="2026-10-25",
            mode="all",
        )
        result = on_day(row, date(2026, 9, 24), 2026, {})[0]["text"]
        self.assertIn("Ауд.: 324", result)
        self.assertIn("Корпус: Львовская", result)

    def test_horizontal_merge_copies_lecture(self):
        w = load_workbook(FIX / "main.xlsx", data_only=True)
        s = w["1 курс"]
        s.merge_cells("M13:S13")
        s["M13"] = "05.10. Общая лекция"
        s["T13"] = "999"
        s["U13"] = "БП"
        out = io.BytesIO()
        w.save(out)
        w.close()
        b = parse_book(out.getvalue())
        for name in ["26КНТ-3", "26КНТ-4"]:
            g = next(g for g in b["groups"].values() if g["name"] == name)
            items = [v for v in plan(b, g) if v["date"] == "2026-10-05" and "Общая лекция" in v["text"]]
            self.assertEqual(len(items), 1)
            self.assertIn("Ауд.: 999", items[0]["text"])

    def test_bad_source_fails(self):
        with self.assertRaises(Exception):
            parse_book(b"<html>sign in</html>")
        with self.assertRaises(ParseError):
            plan(self.book, self.group, self.en1, ["missing"])

    def test_diff_is_semantic_and_has_all_changes(self):
        old = [dict(date="2026-10-05", time="08:00-09:20", text="Математика\nАуд.: 301")]
        new = [dict(old[0], text="Математика\nАуд.: 302")]
        text = "\n".join(changes(old, new, date(2026, 10, 3)))
        self.assertIn("Удалено", text)
        self.assertIn("Добавлено", text)
        self.assertIn("301", text)
        self.assertIn("302", text)
        self.assertEqual(changes(old, old, date(2026, 10, 3)), [])
        self.assertEqual(changes(old, [], date(2026, 10, 6)), [])

    def test_telegram_message_limit(self):
        text = "Предмет\n" * 1800
        values = split(text)
        self.assertGreater(len(values), 1)
        self.assertTrue(all(len(v) <= 3500 for v in values))


