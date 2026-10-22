import asyncio
import copy
import hashlib
import io
import re
import ssl
from datetime import date, datetime, timedelta

import aiohttp
import certifi
import pdfplumber
from openpyxl import load_workbook

from config import SCHEDULE_URL, TZ
from logger import logger
from sources import ParseError, SourceUnavailable, discover_modules, module_number
import database

refs = []
source_index = {}
PARSER_VERSION = 2
lock = asyncio.Lock()
http = None
DAY = {
    "пн": 0,
    "вт": 1,
    "ср": 2,
    "чт": 3,
    "пт": 4,
    "сб": 5,
    "вс": 6,
    "понедельник": 0,
    "вторник": 1,
    "среда": 2,
    "четверг": 3,
    "пятница": 4,
    "суббота": 5,
    "воскресенье": 6,
}
MONTH = {
    "январь": 1,
    "февраль": 2,
    "март": 3,
    "апрель": 4,
    "май": 5,
    "июнь": 6,
    "июль": 7,
    "август": 8,
    "сентябрь": 9,
    "октябрь": 10,
    "ноябрь": 11,
    "декабрь": 12,
}
# Время 08:00-09:20 и даты вроде 07.09 или 07.09-19.10.
TIME = re.compile(r"(\d{1,2}[:.]\d{2})\s*[-–—]\s*(\d{1,2}[:.]\d{2})")
STAMP = r"(?:0?[1-9]|[12]\d|3[01])\.(?:0[1-9]|1[0-2])(?:\.\d{4})?\.?"
DATES = re.compile(r"(?<!\d)(" + STAMP + r")(?!\d)")
SPAN = re.compile(r"(" + STAMP + r")\s*(?:[-–—]|по|до)\s*(" + STAMP + r")", re.I)
LEVEL = re.compile(r"НАЧАЛЬНЫЙ|БАЗОВЫЙ|ОСНОВНОЙ|ПРОДВИНУТЫЙ|ПРАКТИЧЕСКИЙ", re.I)
# Легенда основной таблицы: нужна также для кэша старой версии без buildings.
DEFAULT_BUILDINGS = {
    "БП": "Большая Печёрская, 25/12",
    "БПОК": "Большая Покровская, 22",
    "К": "Костина, 2Б",
    "Л": "Львовская, 1В",
    "Р": "Родионова, 136",
    "С": "Сормовское шоссе, 30",
}
DATE_LIST = r"(?<!\d)" + STAMP + r"(?:\s*(?:[,;]|[-–—]|по|до|to|и)\s*" + STAMP + r")*(?!\d)"
CANCEL = r"(?:отмен\w*|отм\.?(?![а-я])|кроме|не\s+(?:будет|провод\w*)(?:\s+заняти\w*)?)"
CANCEL_NOTE = re.compile(
    r"(?:" + CANCEL + r"\s*[:–—-]?\s*" + DATE_LIST + r"|"
    + DATE_LIST + r"\s*[:–—-]?\s*" + CANCEL + r"(?:\s+заняти\w*)?)", re.I,
)
DATE_CHANGE = re.compile(
    DATE_LIST + r"\s*(?:[-–—]\s*)?(?:[вс]\s+\d{1,2}:\d{2}(?:\s*[-–—]\s*\d{1,2}:\d{2})?|"
    r"в\s+(?:ауд\.?\s*)?\d+[А-Яа-я]?(?:\s*\([^)]+\))?|online|онлайн)", re.I,
)
ASSESSMENT_NOTE = re.compile(
    r"^\s*(?P<dates>" + DATE_LIST + r")\s*(?:[-–—:]\s*)?"
    r"(?P<label>контрольная\s+работа\b|к\.?\s*р\.?(?!\w))", re.I,
)
ROOM_NOTE = re.compile(r"\b(?:в\s+)?ауд(?:итория)?\.?\s*:?\s*(\d+[А-Яа-я]?)", re.I)
CODES = {
    "Филология": {"ФИЛ"},
    "Фундаментальная и прикладная лингвистика": {"ФИПЛ"},
    "Математика": {"МТ", "МАГМТ", "ФПМ", "ФМ"},
    "Фундаментальная и прикладная математика": {"ФПМ", "ФМ"},
    "Дизайн": {"НДИЗ", "ДИЗ"},
    "Компьютерные науки и технологии": {"КНТ"},
    "Технологии искусственного и дополненного интеллекта": {"ТИДИ"},
    "Иностранные языки и межкультурная бизнес-коммуникация": {"ИЯМБК"},
    "Международный бакалавриат по бизнесу и экономике": {"МПБЭ", "МББЭ"},
    "Цифровой маркетинг": {"ЦМ"},
    "Юриспруденция": {"Ю"},
    "Правовое обеспечение и защита бизнеса": {"ПОЗБ"},
    "Управление правовыми рисками в бизнесе": {"УПРБ"},
}


async def open_sources():
    global http, refs, source_index
    source_index = database.get_cache("source-index:v1") or {}
    refs = (database.get_cache("catalog") or []) if source_index else []
    ssl_ctx = ssl.create_default_context(cafile=certifi.where())
    http = aiohttp.ClientSession(connector=aiohttp.TCPConnector(ssl=ssl_ctx), timeout=aiohttp.ClientTimeout(total=60))


async def close_sources():
    if http:
        await http.close()


async def download(url, kind):
    headers = {"xlsx": (b"PK",), "pdf": (b"%PDF",), "html": (b"<!doctype html", b"<html")}
    for attempt in range(3):
        try:
            async with http.get(url) as reply:
                reply.raise_for_status()
                data = await reply.read()
            if len(data) > 20_000_000:
                raise ParseError("Слишком большой файл")
            if not data.lstrip().lower().startswith(tuple(value.lower() for value in headers[kind])):
                raise ParseError("Сайт вернул страницу вместо расписания")
            return data
        except (aiohttp.ClientError, asyncio.TimeoutError):
            if attempt == 2:
                raise
            await asyncio.sleep(2 ** attempt)


async def get_index():
    global refs, source_index
    html = (await download(SCHEDULE_URL, "html")).decode("utf-8")
    modules = discover_modules(html, SCHEDULE_URL)
    rows, parsed = [], {}
    for number, item in modules.items():
        if not item["url"]:
            continue
        try:
            if item["url"] not in parsed:
                parsed[item["url"]] = await asyncio.to_thread(catalog, await download(item["url"], "pdf"))
            item["refs"] = parsed[item["url"]]
            rows.extend(item["refs"])
        except Exception:
            logger.exception("Не удалось загрузить PDF модуля %s", number)
            item["error"] = f"Не удалось загрузить PDF расписания модуля {number}. Повторим проверку позже."
    # Календарь из официальной XLSX содержит даты всех четырёх модулей,
    # сессий и каникул, даже когда PDF следующих модулей ещё не опубликованы.
    calendar_book = None
    for ref in rows:
        if ref["english"]:
            continue
        try:
            candidate = await get_book(ref["book"], fresh=True)
            if {info["module"] for info in candidate.get("calendar_days", {}).values()} == {1, 2, 3, 4} \
                    and candidate["year"] == academic_year(datetime.now(TZ).date()):
                calendar_book = candidate
                break
        except Exception:
            logger.exception("Не удалось прочитать учебный календарь из %s", ref["name"])
    if calendar_book is None:
        raise ParseError("Не удалось получить официальный календарь текущего учебного года")
    index = {"year": calendar_book["year"], "days": calendar_book["calendar_days"], "modules": modules}
    database.save_cache("source-index:v1", index)
    # Стабильные id направлений сохраняют подписки при смене PDF/модуля.
    refs = list({item["id"]: item for item in rows}.values())
    database.save_cache("catalog", refs)
    source_index = index


def academic_year(day):
    return day.year if day.month >= 9 else day.year - 1


def calendar_day(day):
    if not source_index:
        return None
    info = source_index.get("days", {}).get(day.isoformat())
    if source_index.get("year") != academic_year(day) or info is None:
        raise SourceUnavailable(f"Для {day:%d.%m.%Y} ещё нет актуального учебного календаря. Расписание не подтверждено.")
    return info


def get_catalog(day=None):
    day = day or datetime.now(TZ).date()
    info = calendar_day(day)
    if info is None:
        if not refs:
            raise SourceUnavailable("Источники расписания ещё не проверены. Повторите запрос после обновления.")
        return refs
    number = info["module"]
    item = source_index.get("modules", {}).get(str(number), {})
    if item.get("error"):
        raise SourceUnavailable(item["error"])
    if not item.get("url") or not item.get("refs"):
        raise SourceUnavailable(f"Расписание очного обучения для модуля {number} ещё не опубликовано на сайте ВШЭ.")
    return item["refs"]


def source_notice(day):
    try:
        info = calendar_day(day)
        if info and info["kind"] == "session":
            return "По учебному календарю идёт сессия. Расписание экзаменов публикуется отдельно от учебных занятий."
        if info and info["kind"] == "vacation":
            return "По учебному календарю каникулы."
        get_catalog(day)
    except SourceUnavailable as error:
        return str(error)
    return None


async def get_book(bid, english=False, fresh=False):
    cache_key = "book:" + bid
    if english:
        cache_key = "en:" + bid
    if not fresh:
        saved = database.get_cache(cache_key)
        if saved and saved.get("parser_version") == PARSER_VERSION:
            return saved
    url = f"https://docs.google.com/spreadsheets/d/{bid}/export?format=xlsx"
    data = await download(url, "xlsx")
    book = await asyncio.to_thread(parse_book, data, english)
    database.save_cache(cache_key, book)
    return book


def get_program(pid, day=None):
    for item in get_catalog(day):
        if item["id"] == pid:
            return item
    return None


def get_english(course, day=None):
    for item in get_catalog(day):
        if item["english"] and item["name"] == f"{course} курс":
            return item
    return None


async def read_once(bid, english, fresh, seen):
    # За одну проверку общую таблицу скачиваем один раз.
    name = (bid, english)
    if name not in seen:
        try:
            seen[name] = await get_book(bid, english, fresh=fresh)
        except Exception as e:
            seen[name] = e
    result = seen[name]
    if isinstance(result, Exception):
        raise result
    return result


async def get_data(user, fresh=False, seen=None, day=None):
    if seen is None:
        seen = {}
    day = day or datetime.now(TZ).date()
    ref = get_program(user["program"], day)
    if not ref:
        raise ParseError("Направление пропало из указателя. Откройте /settings.")
    book = await read_once(ref["book"], False, fresh, seen)
    group = book["groups"].get(user["group_id"])
    if not group:
        raise ParseError("Учебная группа пропала из таблицы. Откройте /settings.")
    validate_schedule(book, group, day)

    en = None
    if user["en"]:
        ref = get_english(group["course"], day)
        if not ref:
            raise ParseError("Не найдено расписание английского. Откройте /english.")
        en = await read_once(ref["book"], True, fresh, seen)
        validate_english(en, book, user["en"], day)
    return book, group, en


def validate_schedule(book, group, day):
    info = calendar_day(day)
    periods = group.get("periods") or group["rows"]
    if info and (book["year"] != academic_year(day) or info["kind"] == "classes" and not any(
            row.get("module") == info["module"] and row["start"] <= row["end"]
            and academic_year(date.fromisoformat(row["start"])) == book["year"]
            for row in periods)):
        raise SourceUnavailable(f"Таблица группы ещё не содержит расписание на {day:%d.%m.%Y}. Дождитесь публикации источника.")


def validate_english(english, own, ids, day):
    if not source_index:
        return
    aligned = current_english(english, own)
    for gid in ids:
        group = aligned["en"].get(gid)
        if group is None:
            raise ParseError("Выбранная группа английского исчезла. Откройте /english.")
        try:
            validate_schedule(aligned, group, day)
        except SourceUnavailable:
            raise SourceUnavailable(f"Таблица английского ещё не содержит расписание на {day:%d.%m.%Y}.") from None


async def check_sources():
    await get_index()
    books = {}
    all_refs = [ref for item in source_index.get("modules", {}).values() for ref in item.get("refs", [])] if source_index else refs
    for ref in all_refs:
        books[ref["book"]] = ref
    for bid, ref in books.items():
        book = await get_book(bid, ref["english"], fresh=True)
        groups = book["groups"]
        if ref["english"]:
            groups = book["en"]
        for group in groups.values():
            if ref["english"]:
                plan(book, {"rows": []}, book, [group["id"]])
            else:
                plan(book, group)
        print(f"OK: {ref['name']} — {len(groups)} групп")
    print(f"Проверено таблиц: {len(books)}")


def parse_book(data, english=False):
    book = load_workbook(io.BytesIO(data), data_only=True)
    try:
        year, weeks, calendar_days = calendar(book, details=True)
        buildings = read_buildings(book)
        if not english and "Расписание" in book.sheetnames:
            # У дизайна все курсы на одном листе.
            for row in book["Расписание"].iter_rows(max_row=5):
                for cell in row:
                    if re.search(r"гр\.\s*\d{2}ндиз", clean(cell.value), re.I):
                        result = parse_design(book, year, weeks)
                        result["buildings"] = buildings
                        result.update(calendar_days=calendar_days, parser_version=PARSER_VERSION)
                        return result

        groups = {}
        en = {}
        for sheet in book:
            match = re.fullmatch(r"(\d)\s*курс(?:\s+.*)?", sheet.title.strip(), re.I)
            if not match or sheet.sheet_state != "visible":
                continue
            course = int(match[1])
            for first, last in sheet_sections(sheet):
                if english:
                    merge_groups(en, read_english(sheet, course, year, first, last))
                else:
                    merge_groups(groups, read_groups(sheet, course, year, first, last))
        if (english and not en) or (not english and not groups):
            raise ParseError("Не найдено ни одной группы")
        if not english:
            has_lessons = False
            for group in groups.values():
                if group["rows"] or group["refs"]:
                    has_lessons = True
                    break
            if not has_lessons:
                raise ParseError("Таблица содержит группы, но не содержит занятий")
        return {"year": year, "weeks": weeks, "groups": groups, "en": en, "buildings": buildings,
                "calendar_days": calendar_days, "parser_version": PARSER_VERSION}
    finally:
        book.close()


def catalog(data):
    # В PDF только список направлений и ссылки на их таблицы.
    result = []
    level = ""
    faculty = ""
    english = False

    with pdfplumber.open(io.BytesIO(data)) as pdf:
        for page in pdf.pages:
            for line in page.extract_text_lines():
                text = line["text"].strip()
                if text in ["БАКАЛАВРИАТ", "МАГИСТРАТУРА"]:
                    level = text
                    english = False
                elif text.startswith("Факультет"):
                    faculty = text
                    english = False
                elif text == "АНГЛИЙСКИЙ ЯЗЫК":
                    english = True
                    faculty = ""

                for link in page.hyperlinks:
                    url = link.get("uri", "")
                    match = re.search(r"^https://docs.google.com/spreadsheets/d/([\w-]+)", url)
                    if not match:
                        continue
                    if not link["top"] <= line["top"] + 4 <= link["bottom"]:
                        continue
                    result.append({
                        "id": key(level, faculty, text),
                        "name": text,
                        "level": level,
                        "faculty": faculty,
                        "english": english,
                        "book": match[1],
                        "url": url,
                    })

    has_english = False
    for item in result:
        if item["english"]:
            has_english = True
            break
    if not result or not has_english:
        raise ParseError("В PDF не найдены направления или ссылки на английский язык")
    return result


def merge_groups(target, added):
    for gid, group in added.items():
        if gid not in target:
            target[gid] = group
            continue
        target[gid]["rows"].extend(group["rows"])
        if group.get("periods"):
            target[gid].setdefault("periods", []).extend(group["periods"])
        if "refs" in group:
            target[gid]["refs"] |= group["refs"]
        for teacher in group.get("teachers", []):
            if teacher not in target[gid]["teachers"]:
                target[gid]["teachers"].append(teacher)


def sheet_sections(sheet):
    headers = []
    for row in sheet:
        values = {clean(cell.value).casefold() for cell in row}
        if {"день", "время"} <= values:
            headers.append(row[0].row)
    if not headers:
        return [(1, sheet.max_row)]
    return [(1 if i == 0 else row, headers[i + 1] - 1 if i + 1 < len(headers) else sheet.max_row)
            for i, row in enumerate(headers)]


def read_groups(sheet, course, year, first_row=1, last_row=None):
    g = read_grid(sheet)
    layout = slots(g, first_row, last_row)
    if layout is None:
        raise ParseError(f"{sheet.title}: отсутствуют День/Время")
    head, rows = layout
    first = head + 4
    if rows:
        first = rows[0][0]

    headers = []
    for row in sheet.iter_rows(min_row=head, max_row=first - 1, min_col=4):
        for cell in row:
            name = clean(cell.value)
            match = re.fullmatch(r"\d{2}\s*[-–]?\s*[А-Яа-яЁёA-Za-z]+[\w\s-]*", name)
            if match and "модул" not in name:
                headers.append((cell.row, cell.column, name))
    if not headers:
        raise ParseError(f"{sheet.title}: не найдены учебные группы")

    groups = {}
    for hr, col, name in headers:
        gid = key(f"{course} курс", norm(name))
        start, end = period(g, col, hr, year)
        module = module_number(above(g, hr, col, r"модул"))
        room_col = None
        place_col = None
        if grid_text(g, hr, col + 1).lower().startswith("ауд"):
            room_col = col + 1
        if grid_text(g, hr, col + 2).lower() == "корпус":
            place_col = col + 2
        group = {
            "id": gid, "name": name, "code": group_code(name),
            "course": course, "sheet": sheet.title, "rows": [], "refs": False,
            "periods": [{"start": start, "end": end, "module": module}],
        }
        for r, day, time in rows:
            text = grid_text(g, r, col)
            if not text or re.fullmatch(r"[-–—\s]+", text):
                continue
            if "занятия проводятся по расписанию" in text.lower() and "английск" in text.lower():
                group["refs"] = True
                continue

            # У общей лекции аудитория записана после последней группы.
            rc = room_col
            pc = place_col
            area = g["merges"].get((r, col))
            if area and (area.max_col > col or area.min_col < col):
                if room_col and area.max_col < sheet.max_column:
                    rc = area.max_col + 1
                    pc = area.max_col + 2
            room = ""
            place = ""
            if rc:
                room = grid_text(g, r, rc)
            if pc:
                place = grid_text(g, r, pc)
            group["rows"].append({
                "day": day, "time": time, "text": text,
                "room": room, "place": place,
                "start": start, "end": end, "mode": "all",
                "module": module,
            })
        merge_groups(groups, {gid: group})
    return groups


def read_english(sheet, course, year, first_row=1, last_row=None):
    g = read_grid(sheet)
    layout = slots(g, first_row, last_row)
    if layout is None:
        raise ParseError(f"{sheet.title}: отсутствуют День/Время")
    head, rows = layout
    first = head + 4
    if rows:
        first = rows[0][0]

    headers = []
    for row in sheet.iter_rows(min_row=first_row, max_row=first - 1):
        for cell in row:
            if clean(cell.value).lower() == "группа":
                headers.append((cell.row, cell.column))

    groups = {}
    for hr, col in headers:
        faculty = above(g, hr, col, r"^Факультет")
        title = above(g, hr, col, r"ВЕРХНЯЯ НЕДЕЛЯ|НИЖНЯЯ НЕДЕЛЯ")
        mode = "all"
        if "ВЕРХНЯЯ" in title:
            mode = "upper"
        elif "НИЖНЯЯ" in title:
            mode = "lower"
        start, end = period(g, col, hr, year)
        module = module_number(above(g, hr, col, r"модул"))
        with_subject = grid_text(g, hr, col - 1).lower() == "дисциплина"

        for r, day, time in rows:
            for group, subject, teacher, room, place in en_rows(g, r, col, with_subject):
                match = re.search(r"(?:гр(?:уппа)?\.?\s*|\bр\s*)(\d+)", group, re.I)
                num = 0
                if match:
                    num = int(match[1])
                code = ""
                code_match = re.search(r"\(([^)]+)\)", group)
                if code_match:
                    code = re.sub(r"^\d+", "", norm(code_match[1]))
                level = "ОСНОВНОЙ"
                if mode == "lower":
                    level = "МЕЖДУНАРОДНЫЕ ЭКЗАМЕНЫ"
                level_match = LEVEL.search(subject)
                if level_match:
                    level = level_match[0].upper()
                kind = "АЯ"
                if with_subject and "ПРАКТИЧЕСК" in subject.upper():
                    kind = "ПРАКТИЧЕСКИЙ"
                if not match:
                    kind = group

                # Один номер бывает у разных направлений и уровней.
                gid = key(f"{course} курс", faculty, str(num), code, level, mode, kind)
                if gid not in groups:
                    groups[gid] = {
                        "id": gid, "num": num, "code": code, "level": level,
                        "mode": mode, "faculty": faculty, "course": course,
                        "teachers": [], "rows": [],
                    }
                item = groups[gid]
                name = re.split(r"\s*(?:" + STAMP + r"|отм\.?|с\s+\d)", teacher, maxsplit=1, flags=re.I)[0].strip()
                if name and name not in item["teachers"]:
                    item["teachers"].append(name)
                item["rows"].append({
                    "day": day, "time": time, "text": subject,
                    "teacher": teacher, "group": group, "room": room, "place": place,
                    "start": start, "end": end, "mode": mode,
                    "module": module,
                })
    return groups


def parse_design(book, year, weeks):
    # У дизайна четыре курса стоят рядом, поэтому читаем их отдельно.
    sheet = book["Расписание"]
    g = read_grid(sheet)
    groups = {}
    for row in sheet.iter_rows(max_row=5):
        for cell in row:
            match = re.search(r"гр\.\s*(\d{2}ндиз-\d+)", clean(cell.value), re.I)
            if not match:
                continue
            col = cell.column
            hr = cell.row
            title = above(g, hr, col, r"\d\s*КУРС")
            course = int(re.search(r"\d", title)[0])
            meta = read_grid(book[f"{course} курс"])
            head, _ = slots(meta)
            start, end = period(meta, 5, head, year)
            module = module_number(above(meta, head, 5, r"модул"))
            name = match[1].upper()
            gid = key(f"{course} курс", norm(name))
            group = {
                "id": gid, "name": name, "code": "НДИЗ", "course": course,
                "sheet": sheet.title, "rows": [], "refs": False, "note": clean(cell.value),
                "periods": [{"start": start, "end": end, "module": module}],
            }
            time_col = None
            for c in range(col - 1, 0, -1):
                if grid_text(g, hr, c).lower() == "время":
                    time_col = c
                    break
            if time_col is None:
                raise ParseError(f"{sheet.title}: не найдена колонка времени")

            day = None
            for r in range(hr + 1, sheet.max_row + 1):
                text = grid_text(g, r, 1).lower()
                if text in DAY:
                    day = DAY[text]
                match = TIME.fullmatch(grid_text(g, r, time_col))
                text = grid_text(g, r, col)
                if not match or day is None or not text:
                    continue
                if re.search(r"(?:ГРУППА|ГР\.)\s*\d+.*(?:английск|АЯ)", text, re.I):
                    group["refs"] = True
                    continue
                area = g["merges"].get((r, col))
                room = ""
                if not area or area.max_col == col:
                    room = grid_text(g, r, col + 1)
                time = match[1].zfill(5) + "-" + match[2].zfill(5)
                group["rows"].append({
                    "day": day, "time": time, "text": text, "room": room, "place": "",
                    "start": start, "end": end, "mode": "all",
                    "module": module,
                })
            groups[gid] = group
    if not groups:
        raise ParseError("Не найдены группы дизайна")
    return {"year": year, "weeks": weeks, "groups": groups, "en": {}}


def clean(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).replace("\xa0", " ").replace("\r", "").strip()


def norm(value):
    text = clean(value).upper()
    text = re.sub(r"\s+", "", text)
    text = text.replace("Ё", "Е")
    text = text.replace("–", "-").replace("—", "-")
    return text


def building_key(value):
    text = re.sub(r"[^А-ЯЁA-Z0-9]", "", clean(value).upper()).replace("Ё", "Е")
    return {"P": "Р", "K": "К", "C": "С"}.get(text, text)


def building_address(value):
    text = re.sub(r"^корпус\s+на\s+", "", clean(value), flags=re.I)
    text = re.sub(r"^ул\.?\s*", "", text, flags=re.I)
    text = re.sub(r"\bБ\.\s*", "Большая ", text)
    text = re.sub(r"ской\b", "ская", text)
    text = re.sub(r"Сормовском\s+ш\.?", "Сормовское шоссе", text, flags=re.I)
    return re.sub(r"\s+", " ", text).strip()


def merge_buildings(*sources):
    result = {}
    for source in sources:
        for code, address in source.items():
            code = building_key(code)
            # Легенда английского иногда содержит только улицу. Сохраняем дом
            # из основной таблицы, даже если она передана первой.
            if code not in result or re.search(r"\d", address) or not re.search(r"\d", result[code]):
                result[code] = address
    return result


def read_buildings(book):
    result = {}
    for sheet in book:
        for row in sheet:
            for cell in row:
                text = clean(cell.value)
                if re.match(r"обозначения\s+корпусов\s*:", text, re.I):
                    text = text.split(":", 1)[1]
                    for item in re.split(r"\s*/\s*|\n", text):
                        match = re.fullmatch(r"\s*(.+?)\s*[-–—]\s*([А-ЯЁA-Zа-яёa-z.]{1,8})\s*", item)
                        if match:
                            result = merge_buildings(result, {match[2]: building_address(match[1])})
                elif re.match(r"корпус\s+на\s+", text, re.I):
                    # В обычном расписании адрес и сокращение стоят рядом.
                    for offset in [1, 2, 3, -1]:
                        col = cell.column + offset
                        if not 1 <= col <= sheet.max_column:
                            continue
                        code = clean(sheet.cell(cell.row, col).value)
                        if re.fullmatch(r"[А-ЯЁA-Zа-яёa-z.]{1,8}", code):
                            result = merge_buildings(result, {code: building_address(text)})
                            break
    return result


def resolve_building(value, buildings):
    aliases = {}
    for code, address in buildings.items():
        name = re.split(r",?\s*\d", address, maxsplit=1)[0]
        for alias in [code, address, name, name.replace("Большая ", "Б")]:
            aliases[building_key(alias)] = address

    result = []
    for line in clean(value).split("\n"):
        name = building_key(line)
        address = aliases.get(name)
        if address is None and len(name) >= 3:
            # В строках встречаются Б.Покр / БПокр и Б.Печ / БПеч,
            # хотя в легенде используются БПок и БП.
            matches = {address for alias, address in aliases.items() if alias.startswith(name)}
            if len(matches) == 1:
                address = matches.pop()
        result.append(address or line)
    return "\n".join(result)


def expand_building(lesson, buildings=None):
    if buildings is None:
        buildings = DEFAULT_BUILDINGS
    result = dict(lesson)
    before, mark, place = result["text"].rpartition("\nКорпус: ")
    if mark:
        result["building"] = resolve_building(place, buildings)
        result["text"] = before + mark + result["building"]
    elif result.get("building"):
        result["building"] = resolve_building(result["building"], buildings)
    return result


def expand_plan_buildings(lessons, buildings=None):
    return [expand_building(lesson, buildings) for lesson in lessons]


def key(*parts):
    return hashlib.sha256("|".join(parts).encode()).hexdigest()[:20]


def read_grid(sheet):
    # У объединённой ячейки текст лежит только в левом верхнем углу.
    merges = {}
    for area in sheet.merged_cells.ranges:
        for r in range(area.min_row, area.max_row + 1):
            for c in range(area.min_col, area.max_col + 1):
                merges[r, c] = area
    return {"sheet": sheet, "merges": merges}


def grid_cell(g, r, c):
    area = g["merges"].get((r, c))
    if area:
        r, c = area.min_row, area.min_col
    return g["sheet"].cell(r, c)


def grid_text(g, r, c):
    return clean(grid_cell(g, r, c).value)


def above(g, row, col, pattern):
    for r in range(row - 1, 0, -1):
        value = grid_text(g, r, col)
        if re.search(pattern, value, re.I):
            return value
    return ""


def color(cell):
    if cell.fill.patternType == "solid":
        return str(cell.fill.fgColor)
    return ""


def calendar(book, details=False):
    sheet = None
    for s in book:
        if "календар" in s.title.lower():
            sheet = s
            break
    if sheet is None:
        raise ParseError("Не найден календарь с учебным годом")

    match = re.search(r"(20\d{2})\s*[-–/]\s*20\d{2}", clean(sheet["A1"].value))
    if not match:
        raise ParseError("Не найден учебный год")
    year = int(match[1])

    upper = None
    for row in sheet:
        for cell in row:
            if clean(cell.value).lower() == "верхняя неделя":
                upper = color(cell)
                break
        if upper is not None:
            break

    g = read_grid(sheet)
    weeks = {}
    days = {}
    module = None
    phases = {clean(cell.value).casefold(): color(cell) for row in sheet for cell in row
              if clean(cell.value).casefold() in {"сессия", "каникулы"}}
    month_row = None
    for row in sheet:
        for cell in row:
            number = module_number(clean(cell.value))
            if number:
                module = number
        has_month = False
        for cell in row:
            if clean(cell.value).lower() in MONTH:
                has_month = True
                break
        if has_month:
            month_row = row[0].row
            continue
        if month_row is None:
            continue

        for cell in row:
            if not isinstance(cell.value, (float, int)):
                continue
            if not 1 <= cell.value <= 31:
                continue

            # Ищем месяц слева: его название может занимать несколько ячеек.
            month = None
            for col in range(cell.column, 0, -1):
                text = grid_text(g, month_row, col).lower()
                if text in MONTH:
                    month = MONTH[text]
                    break
            if month is None:
                continue

            day_year = year
            if month < 9:
                day_year += 1
            try:
                day = date(day_year, month, int(cell.value))
            except ValueError:
                logger.warning("Некорректная дата в календаре: %s", cell.coordinate)
                continue
            mode = "lower"
            if upper and color(cell) == upper:
                mode = "upper"
            weeks[day.isoformat()] = mode
            if module:
                kind = "classes"
                if phases.get("сессия") and color(cell) == phases["сессия"]:
                    kind = "session"
                elif phases.get("каникулы") and color(cell) == phases["каникулы"]:
                    kind = "vacation"
                stamp = day.isoformat()
                info = {"module": module, "kind": kind}
                if stamp in days and days[stamp] != info:
                    raise ParseError(f"Противоречивые модули в календаре на {day}")
                days[stamp] = info
    return (year, weeks, days) if details else (year, weeks)


def as_date(text, year):
    nums = re.findall(r"\d+", text)
    day = int(nums[0])
    month = int(nums[1])
    if len(nums) > 2:
        year = int(nums[2])
    elif month < 9:
        year += 1
    return date(year, month, day)


def period(g, col, row, year):
    title = above(g, row, col, r"модул")
    spans = SPAN.findall(title)
    if not spans:
        raise ParseError(f"{g['sheet'].title}: не найден срок модуля для колонки {col}")
    start = as_date(spans[0][0], year)
    end = as_date(spans[0][1], year)
    return start.isoformat(), end.isoformat()


def slots(g, first_row=1, last_row=None):
    day_col = None
    time_col = None
    head = None
    last_row = last_row or g["sheet"].max_row
    for row in g["sheet"].iter_rows(min_row=first_row, max_row=last_row):
        for cell in row:
            text = clean(cell.value).lower()
            if text == "день":
                day_col = cell.column
            elif text == "время":
                time_col = cell.column
        if day_col and time_col:
            head = row[0].row
            break
    if head is None:
        return None

    result = []
    day = None
    for r in range(head + 1, last_row + 1):
        text = grid_text(g, r, day_col).split("\n")[0].lower().rstrip(".")
        if text in DAY:
            day = DAY[text]
        match = TIME.fullmatch(grid_text(g, r, time_col))
        if not match or day is None:
            continue
        start = match[1].replace(".", ":").zfill(5)
        end = match[2].replace(".", ":").zfill(5)
        result.append((r, day, start + "-" + end))
    return head, result


def group_code(name):
    return re.sub(r"[^А-ЯA-Z]", "", norm(name))


def lines(text):
    result = []
    for line in clean(text).split("\n"):
        line = line.strip()
        if line:
            result.append(line)
    return result


def en_align(text, starts, group_lines, field, required=False):
    values = lines(text)
    count = len(starts)
    if len(values) == count:
        return values
    if not required and len(values) <= 1:
        if not values:
            values = [""]
        return values * count

    # Пустые строки важны: по ним сопоставляем группу с преподавателем.
    raw = clean(text).split("\n")
    while len(raw) < group_lines:
        raw.append("")
    for i in range(group_lines):
        if i not in starts and raw[i].strip():
            raise ParseError(f"{field}: нельзя сопоставить строки английского")

    result = []
    for i in starts:
        result.append(raw[i].strip())
    tail = " / ".join(lines("\n".join(raw[group_lines:])))
    if tail:
        for i in range(count):
            result[i] += "\nПримечание общей ячейки: " + tail
    if required and not all(result):
        raise ParseError(f"{field}: нельзя сопоставить строки английского")
    return result


def row_blocks(values, starts):
    result = []
    for i, start in enumerate(starts):
        end = len(values)
        if i + 1 < len(starts):
            end = starts[i + 1]
        text = " ".join(lines("\n".join(values[start:end])))
        result.append(text)
    return result


def en_rows(g, r, col, with_subject):
    text = grid_text(g, r, col)
    values = text.split("\n")
    starts = []
    for i, line in enumerate(values):
        number = re.search(r"(?:гр(?:уппа)?\.?\s*|(?<![А-Яа-я])р\s*)\d+", line, re.I)
        special = "английский язык для специальных целей" in line.lower()
        if number or special:
            starts.append(i)
    if not starts:
        if text:
            raise ParseError(f"{g['sheet'].title}, {r}:{col}: неизвестный формат английских групп")
        return []

    groups = row_blocks(values, starts)
    teachers = en_align(grid_text(g, r, col + 1), starts, len(values), "Преподаватели", required=True)
    rooms = en_align(grid_text(g, r, col + 2), starts, len(values), "Аудитории")
    places = en_align(grid_text(g, r, col + 3), starts, len(values), "Корпуса")
    subjects = []
    if with_subject:
        raw = grid_text(g, r, col - 1).split("\n")
        starts = []
        for i, line in enumerate(raw):
            if re.search(r"АЯ|ЯДОК|английск", line, re.I):
                starts.append(i)
        subjects = row_blocks(raw, starts)
        if len(subjects) != len(groups):
            raise ParseError(f"{g['sheet'].title}, строка {r}: дисциплины английского не совпадают с группами")
    else:
        for group in groups:
            subjects.append("Английский язык: " + group)
    return zip(groups, subjects, teachers, rooms, places)


def skip_time(match):
    # 11:10-12:30 — время, а 07.09-19.10 может быть периодом занятий.
    if ":" in match[0] or int(match[2][-2:]) > 12:
        return ""
    return match[0]


def swap_date(match):
    return f"{int(match[2]):02}.{int(match[1]):02}"


def date_text(text):
    # В источниках встречаются латинское «c», пробел внутри даты и даты MM/DD.
    text = re.sub(r"(?<=\d)\.\s+(?=\d{2}\b)", ".", text)
    text = re.sub(r"(?<!\d)0\.(\d)/(\d{2})(?!\d)", r"0\1/\2", text)
    text = re.sub(r"(?<!\d)(\d{1,2})/(\d{1,2})(?!\d)", swap_date, text)
    text = re.sub(r"\bfrom\b", "с", text, flags=re.I)
    text = re.sub(r"\bto\b", "по", text, flags=re.I)
    return re.sub(r"\bc(?=\s+\d{1,2}\.)", "с", text, flags=re.I)


def marked_days(text, year):
    result = set()
    for match in DATES.finditer(text):
        result.add(as_date(match[1], year))
    for span in SPAN.finditer(text):
        day = as_date(span[1], year)
        end = as_date(span[2], year)
        while day <= end:
            result.add(day)
            day += timedelta(days=1)
    return result


def date_change(line):
    # Дата изменения времени/места не ограничивает все остальные недели.
    match = DATE_CHANGE.search(line)
    if match and not line[match.end():].strip(" ,;:–—-()"):
        return True
    if not re.search(r"\b(?:в\s*\d|с\s*\d{1,2}:\d{2}|ауд|online|онлайн|корпус|вместо|перенос)", line, re.I):
        return False
    rest = DATES.sub("", line)
    rest = TIME.sub("", rest)
    return len(re.findall(r"[А-Яа-яA-Za-z]+", rest)) <= 4


def date_rules(text, year):
    # Сначала выделяем исключения, чтобы «с 09.09 ..., кроме 23.09» не потеряло начало.
    text = date_text(text)
    allowed = set()
    denied = set()
    restricted = False

    for note in CANCEL_NOTE.finditer(text):
        denied.update(marked_days(note[0], year))
    text = CANCEL_NOTE.sub("", text)

    for line in text.split("\n"):
        if date_change(line) and not ASSESSMENT_NOTE.match(line):
            continue
        line = TIME.sub(skip_time, line)
        marks = list(DATES.finditer(line))
        if not marks:
            continue
        if re.search(CANCEL, line, re.I):
            for match in marks:
                denied.add(as_date(match[1], year))
            continue

        restricted = True
        used = set()
        for span in SPAN.finditer(line):
            start = as_date(span[1], year)
            end = as_date(span[2], year)
            while start <= end:
                allowed.add(start)
                start += timedelta(days=1)
            for i, match in enumerate(marks):
                if span.start() <= match.start() < span.end():
                    used.add(i)

        for i, match in enumerate(marks):
            if i in used:
                continue
            day = as_date(match[1], year)
            before = line[:match.start()]
            if re.search(r"\bс\s*$", before, re.I):
                end = date(year + 1, 8, 31)
                while day <= end:
                    allowed.add(day)
                    day += timedelta(days=1)
            elif re.search(r"\b(?:до|по)\s*$", before, re.I):
                start = date(year, 9, 1)
                while start <= day:
                    allowed.add(start)
                    start += timedelta(days=1)
            else:
                allowed.add(day)
    if not restricted:
        allowed = None  # Даты не указаны — занятие каждую неделю.
    return allowed, denied


def strip_date_notes(line):
    assessment = ASSESSMENT_NOTE.match(line)
    if assessment:
        line = line[:assessment.start("label")] + "Контрольная работа" + line[assessment.end("label"):]
        line = ROOM_NOTE.sub("", line)
    line = CANCEL_NOTE.sub("", line)
    # Убираем смену аудитории/времени вместе с датой, сохраняя имя преподавателя.
    line = DATE_CHANGE.sub("", line)
    line = re.sub(r"\b(?:с|со|до|по)\s+(?=" + STAMP + ")", "", line, flags=re.I)
    line = re.sub(DATE_LIST, "", line, flags=re.I)
    line = re.sub(r"\(\s*\)", "", line)
    return re.sub(r"\s+", " ", line).strip(" ,;:–—-()")


def clean_dates(text, day, year):
    # Правила уже применены к занятию. В сообщении остаются только действующие строки.
    result = []
    for line in date_text(text).split("\n"):
        has_dates = bool(DATES.search(line))
        if not has_dates:
            result.append(line.strip())
            continue
        allowed, _ = date_rules(line, year)
        if allowed is not None and day not in allowed:
            continue
        line = strip_date_notes(line)
        if line:
            result.append(line)
    return "\n".join(value for value in result if value)


def base_location(text):
    result = []
    for line in text.split("\n"):
        mark = DATES.search(date_text(line))
        if mark:
            line = date_text(line)[:mark.start()]
        line = line.strip(" ,;:–—-()")
        if line:
            result.append(line)
    return "\n".join(result)


def lesson_location(room, place, index, count, day, year):
    rooms = lines(room)
    places = lines(place)
    # В соседних ячейках строки соответствуют друг другу:
    # «402 / 30.09 в 203» и «К / БПок». Второй корпус — часть замены.
    paired = len(rooms) == len(places) and any(DATES.search(date_text(v)) for v in rooms)
    if paired:
        pairs = list(zip(rooms, places))
        starts = [i for i, (value, _) in enumerate(pairs) if not DATES.search(date_text(value))]
        if count > 1 and len(starts) == count:
            end = starts[index + 1] if index + 1 < count else len(pairs)
            pairs = pairs[starts[index]:end]
        elif count > 1 and len(pairs) == count and not starts:
            pairs = [pairs[index]]

        base = []
        changes = []
        for room, place in pairs:
            note = date_text(room)
            if not DATES.search(note):
                base.append((room, place))
                continue
            allowed, denied = date_rules(note, year)
            if allowed is None:
                allowed = marked_days(note, year)
            if day in allowed and day not in denied:
                changes.append((room, place))
        selected = changes[-1:] or base
        return "\n".join(v[0] for v in selected), "\n".join(v[1] for v in selected)

    room = rooms[index] if len(rooms) == count else "\n".join(rooms)
    place = places[index] if len(places) == count else "\n".join(places)
    return room, place


def parts(text):
    # В одной ячейке могут быть разные предметы для разных дат.
    result = []
    for block in re.split(r"\n\s*\n+", text):
        if not block.strip():
            continue
        note = re.match(r"^\s*" + STAMP + r"\s*(?:[-–—]\s*)?(?:отмен|отм\b|в\s)", block, re.I)
        if result and (note or ASSESSMENT_NOTE.match(date_text(block))):
            result[-1] += "\n" + block
        else:
            result.append(block.strip())
    return result


def assessment_text(text, day, year):
    regular, selected = [], []
    for line in text.split("\n"):
        note = ASSESSMENT_NOTE.match(date_text(line))
        if note is None:
            regular.append(line)
        elif day in marked_days(note["dates"], year):
            selected.append(line)
    base = "\n".join(regular)
    if not selected:
        return base
    _, denied = date_rules(base, year)
    if day in denied:
        return ""
    # Контрольная наследует предмет, но имеет собственную дату, даже если
    # обычные семинары закончились. Индекс исходного блока сохраняет его корпус.
    context = [strip_date_notes(date_text(line)) for line in regular]
    return "\n".join([*context, *selected])


def on_day(row, day, year, weeks, module=None):
    stamp = day.isoformat()
    if module is not None and row.get("module") is not None and row["module"] != module:
        return []
    if day.weekday() != row["day"]:
        return []
    if not row["start"] <= stamp <= row["end"]:
        return []
    if row["mode"] != "all":
        if stamp not in weeks:
            raise ParseError(f"В календаре нет {day}")
        if weeks[stamp] != row["mode"]:
            return []

    match = re.search(r"со?\s+(\d)\s*модул", row["text"], re.I)
    if match:
        current = row.get("module")
        if current is None:
            raise ParseError("В источнике не указан модуль занятия. Обновите таблицу расписания.")
        if int(match[1]) > current:
            return []
    if re.search(r"нечетным\s+недел", row["text"], re.I) and weeks.get(stamp) != "lower":
        return []

    chunks = parts(row["text"])
    result = []
    for i, text in enumerate(chunks):
        text = assessment_text(text, day, year)
        if not text.strip():
            continue
        allowed, denied = date_rules(text, year)
        notes = "\n".join(row.get(field, "") for field in ["teacher", "group", "room", "place"])
        _, extra_denied = date_rules(notes, year)
        if day in denied or day in extra_denied:
            continue
        if allowed is not None and day not in allowed:
            continue
        outside = False
        for field in ["teacher", "group"]:
            dates, _ = date_rules(row.get(field, ""), year)
            if dates is not None and day not in dates:
                outside = True
        if outside:
            continue
        room, place = lesson_location(
            row.get("room", ""), row.get("place", ""), i, len(chunks), day, year,
        )
        time = row["time"]

        # Проверяем примечания вроде «24.09. в 324 (Львовская)».
        for note in [text, row.get("teacher", ""), room, place]:
            for line in date_text(note).split("\n"):
                dates = []
                for mark in DATES.findall(line):
                    dates.append(as_date(mark, year))
                if day not in dates:
                    continue
                if re.search(r"\bonline\b|онлайн", line, re.I):
                    room = ""
                    place = "Онлайн"
                new_time = TIME.search(line)
                if new_time and re.search(r"\b[вс]\s*\d", line):
                    time = new_time[0].replace(" ", "").replace("–", "-").replace("—", "-")
                start_time = re.search(r"\b[вс]\s+(\d{1,2}:\d{2})(?!\s*[-–—])", line)
                if start_time:
                    start = start_time[1].zfill(5)
                    duration = int(row["time"][6:8]) * 60 + int(row["time"][9:11])
                    duration -= int(row["time"][:2]) * 60 + int(row["time"][3:5])
                    end = int(start[:2]) * 60 + int(start[3:]) + duration
                    time = f"{start}-{end // 60:02}:{end % 60:02}"
                explicit_room = ROOM_NOTE.search(line)
                new_room = explicit_room or re.search(r"\bв\s+(?:ауд\.?\s*)?(\d+[А-Яа-я]?)", line)
                if not new_room:
                    new_room = re.search(STAMP + r"\s*[-–—]\s*(\d+[А-Яа-я]?)\s*\(", line)
                if new_room and (explicit_room or not new_time and not start_time):
                    room = new_room[1]
                    building = re.search(r"\(([^)]+)\)", line)
                    if building:
                        place = building[1]

        info = [clean_dates(text, day, year)]
        if row.get("teacher"):
            info.append(clean_dates(row["teacher"], day, year))
        room = base_location(room)
        place = base_location(place)
        if room:
            info.append("Ауд.: " + room)
        if place:
            info.append("Корпус: " + place)
        result.append({"date": stamp, "time": time, "text": "\n".join(value for value in info if value)})
    return result


def en_options(book, program, group, mode=None):
    codes = set(CODES.get(program["name"], {group["code"]}))
    codes.add(group["code"])
    if group["code"] == "НДИЗ":
        codes.add("ДИЗ")

    faculties = {program["faculty"]}
    for item in book["en"].values():
        if item["course"] == group["course"] and item["code"] in codes:
            faculties.add(item["faculty"])

    result = []
    for item in book["en"].values():
        if item["course"] != group["course"]:
            continue
        if mode and item["mode"] != mode:
            continue
        if item["code"]:
            if item["code"] not in codes:
                continue
        elif item["faculty"] not in faculties:
            continue
        result.append(item)
    return result


def current_english(english, own):
    # В английском второго курса пока старый календарь. Берём год и недели основного.
    if english["year"] == own["year"]:
        return english
    result = copy.deepcopy(english)
    delta = own["year"] - english["year"]
    for group in result["en"].values():
        for row in group["rows"]:
            for field in ["start", "end"]:
                day = date.fromisoformat(row[field])
                day = day.replace(year=day.year + delta)
                row[field] = day.isoformat()
    result["year"] = own["year"]
    result["weeks"] = own["weeks"]
    result["calendar_days"] = own.get("calendar_days", {})
    return result


def lesson_order(item):
    return item["date"], item["time"], item["text"]


def plan(book, group, english=None, ids=()):
    rows = []
    for row in group["rows"]:
        rows.append((row, book))
    if english:
        english = current_english(english, book)
        for gid in ids:
            if gid not in english["en"]:
                raise ParseError("Выбранная группа английского исчезла. Откройте /english.")
            for row in english["en"][gid]["rows"]:
                rows.append((row, english))

    # Считаем занятия для каждого дня модуля. Одинаковые пары не повторяем.
    result = {}
    for row, source in rows:
        buildings = merge_buildings(DEFAULT_BUILDINGS, book.get("buildings", {}), source.get("buildings", {}))
        day = date.fromisoformat(row["start"])
        end = date.fromisoformat(row["end"])
        if (end - day).days > 370:
            raise ParseError("Некорректный срок модуля")
        while day <= end:
            module = source.get("calendar_days", {}).get(day.isoformat(), {}).get("module")
            for lesson in on_day(row, day, source["year"], source["weeks"], module=module):
                lesson = expand_building(lesson, buildings)
                result[lesson_order(lesson)] = lesson
            day += timedelta(days=1)
    return sorted(result.values(), key=lesson_order)
